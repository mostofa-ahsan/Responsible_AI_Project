"""S3 key-fact decomposition of reference answers and S4 claim decomposition of model answers (local judge
model, guided JSON, cached). No API calls.

keyfacts  1-5 atomic, self-contained facts per reference answer, taken ONLY from the reference (the source
          evidence is passed as context for disambiguation). Shared by every system.
          -> data/eval_keyfacts/<split>.jsonl, results/local_eval/keyfact_spotcheck.xlsx (60 random items)
claims    atomic factual claims (<= 12) asserted by each model answer (15 systems x eval subset), processed in
          a random interleaved order so a time-boxed partial run is a random subset.
          -> results/local_eval/per_item/claims.jsonl

Usage:
    python src/local_facts.py keyfacts
    python src/local_facts.py claims
"""

import argparse
import json
import random

from local_common import (CACHE, JUDGES, KEYFACTS, OUT, PER_ITEM, SPLITS, deadline_from_env, items, log, mark,
                          parse_json, preds, read_jsonl, run_worker, shuffled, sys_id, systems, write_jsonl)

KF_SYSTEM = """You decompose a reference answer into atomic key facts that will be used to check other answers.

Rules:
- Write 1 to 5 facts. Each fact is one short declarative sentence that can be checked on its own.
- Make every fact self-contained: resolve pronouns and vague references using the question (name the subject).
- Keep numbers, dates, names, lists and technical terms exactly as the reference states them.
- Take facts ONLY from the reference answer. The source evidence is context for understanding; never add facts from it.
- Skip framing that carries no checkable content. Merge trivial fragments; split compound statements.
Respond with JSON: {"facts": [...]}."""

KF_PROMPT = """Question: {question}

Reference answer: {reference}

Source evidence (context only, do not take facts from it): {evidence}"""

CL_SYSTEM = """You decompose a model's answer into the atomic factual claims it asserts.

Rules:
- List at most 12 claims, each one short, self-contained declarative sentence (resolve pronouns using the question).
- Include only content the answer actually asserts as fact; skip hedges, advice phrased generically, headings,
  restatements of the question and pure framing.
- Keep names, numbers and specific terms exactly as written. If the answer makes more than 12 claims, keep the
  12 most specific ones. If it makes no factual claim, return an empty list.
Respond with JSON: {"claims": [...]}."""

CL_PROMPT = """Question: {question}

Model answer: {answer}"""


def judge():
    p = OUT / "judge_choice.json"
    j = json.loads(p.read_text())["judge"] if p.exists() else "A"
    return j, JUDGES[j]


def extra(jinfo):
    return ["--mistral3"] if jinfo["mistral3"] else []


def cmd_keyfacts(args):
    lg = log()
    deadline = deadline_from_env(30)
    j, ji = judge()
    its = items()
    schema = {"type": "object", "properties": {"facts": {"type": "array", "minItems": 1, "maxItems": 5,
                                                         "items": {"type": "string", "maxLength": 300}}},
              "required": ["facts"], "additionalProperties": False}
    cp = CACHE / f"keyfacts__{ji['name']}.jsonl"
    rows = []
    for s in SPLITS:
        for q, r in its[s].items():
            ev = (r.get("evidence") or "")[:2000]
            rows.append({"id": f"{s}|{q}", "messages": [
                {"role": "system", "content": KF_SYSTEM},
                {"role": "user", "content": KF_PROMPT.format(question=r["question"], reference=r["answer"],
                                                             evidence=ev or "(none)")}]})
    lg.info(f"[S3] key facts for {len(rows)} reference answers with {ji['name']}")
    run_worker(ji["model"], rows, cp, schema=schema, max_tokens=500, deadline=deadline, extra=extra(ji), label="S3")
    got = {}
    for r in read_jsonl(cp):
        g = parse_json(r.get("text", ""))
        if g and isinstance(g.get("facts"), list):
            facts = [f.strip() for f in g["facts"] if isinstance(f, str) and f.strip()][:5]
            if facts:
                got[r["id"]] = facts
    n = 0
    for s in SPLITS:
        out = []
        for q, r in its[s].items():
            f = got.get(f"{s}|{q}")
            if f:
                out.append({"qa_id": q, "split": s, "question": r["question"], "reference": r["answer"], "facts": f,
                            "decomposer": ji["name"]})
        write_jsonl(KEYFACTS / f"{s}.jsonl", out)
        n += len(out)
    lg.info(f"[S3] key facts: {n}/{len(rows)} references decomposed; "
            f"mean {sum(len(v) for v in got.values()) / max(len(got), 1):.2f} facts each")
    spotcheck()
    mark("S3_keyfacts_info", {"decomposed": n, "total": len(rows)})
    if n < 0.95 * len(rows):
        raise SystemExit(f"S3 incomplete: {n}/{len(rows)}")


def spotcheck(n=60, seed=7):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    allf = [r for s in SPLITS for r in read_jsonl(KEYFACTS / f"{s}.jsonl")]
    if not allf:
        return
    sel = random.Random(seed).sample(allf, min(n, len(allf)))
    wb = Workbook()
    ws = wb.active
    ws.title = "keyfacts"
    hdr = ["#", "split", "qa_id", "question", "reference answer", "key facts", "ok? (y/n)", "notes"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, r in enumerate(sel, 1):
        ws.append([i, r["split"], r["qa_id"], r["question"], r["reference"],
                   "\n".join(f"{k}. {f}" for k, f in enumerate(r["facts"], 1)), "", ""])
    for col, w in zip("ABCDEFGH", [4, 16, 30, 50, 60, 70, 10, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(OUT / "keyfact_spotcheck.xlsx")


def cmd_claims(args):
    lg = log()
    deadline = deadline_from_env(45)
    j, ji = judge()
    its = items()
    schema = {"type": "object", "properties": {"claims": {"type": "array", "maxItems": 12,
                                                          "items": {"type": "string", "maxLength": 300}}},
              "required": ["claims"], "additionalProperties": False}
    cp = CACHE / f"claims__{ji['name']}.jsonl"
    rows = []
    for fam, var in systems():
        sid = sys_id(fam, var)
        for s in SPLITS:
            p = preds(fam, var, s)
            for q, r in its[s].items():
                if q in p:
                    rows.append({"id": f"{s}|{q}|{sid}", "messages": [
                        {"role": "system", "content": CL_SYSTEM},
                        {"role": "user", "content": CL_PROMPT.format(question=r["question"], answer=p[q])}]})
    rows = shuffled(rows, seed=11)
    lg.info(f"[S4] claim decomposition for {len(rows)} answers with {ji['name']} (random order; time-boxed)")
    run_worker(ji["model"], rows, cp, schema=schema, max_tokens=900, deadline=deadline, extra=extra(ji) + ["--chunk", "600"], label="S4")
    out = []
    for r in read_jsonl(cp):
        g = parse_json(r.get("text", ""))
        if g is None or not isinstance(g.get("claims"), list):
            continue
        s, q, sid = r["id"].split("|")
        out.append({"split": s, "qa_id": q, "system": sid,
                    "claims": [c.strip() for c in g["claims"] if isinstance(c, str) and c.strip()][:12],
                    "decomposer": ji["name"]})
    write_jsonl(PER_ITEM / "claims.jsonl", out)
    lg.info(f"[S4] claims: {len(out)}/{len(rows)} answers decomposed")
    mark("S4_claims_info", {"decomposed": len(out), "total": len(rows)})


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["keyfacts", "claims", "spotcheck"])
    args = ap.parse_args()
    {"keyfacts": cmd_keyfacts, "claims": cmd_claims, "spotcheck": lambda a: spotcheck()}[args.cmd](args)


if __name__ == "__main__":
    main()

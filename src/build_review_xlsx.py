"""Build a human-review workbook (Instructions / Review / Agreement) for a QA run.

The Review sheet lists every pair of data/qa_pairs/<run>.jsonl with dropdowns for
human_verdict (accept/reject) and reject_category; a stratified priority sample comes first
(priority = TRUE, highlighted): repaired pairs that passed, some rejected pairs, and passed
pairs spread over q_type, chunk origin and difficulty. The Agreement sheet computes progress,
the confusion matrix against the automatic filters and Cohen's kappa with formulas.

Usage:
    python src/build_review_xlsx.py --run pilot_v2 [--out data/qa_pairs/pilot_review_v2.xlsx]
        [--n-priority 40] [--n-repaired 10] [--n-rejected 4] [--seed 7]
"""

import argparse
import random
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from qa_common import read_jsonl, run_paths
from utils import load_config, repo_path

HEADER_FILL = PatternFill("solid", fgColor="FF1F3864")
VERDICT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
CATEGORIES = ["incorrect", "unsupported", "not_standalone", "low_value", "unclear", "other"]

# (header, width, value getter)
COLUMNS = [
    ("#", 5, None), ("priority", 9, None), ("human_verdict", 14, None), ("reject_category", 18, None),
    ("human_notes", 30, None),
    ("question", 45, lambda r: r["question"]),
    ("answer", 55, lambda r: r["answer"]),
    ("citation", 28, lambda r: citation_text(r["citation"])),
    ("evidence", 60, lambda r: r["evidence"]),
    ("passed_filters", 10, lambda r: r["passed_filters"]),
    ("reject_reason", 20, lambda r: r["reject_reason"]),
    ("repaired", 10, lambda r: bool(r["repair"])),
    ("q_type", 12, lambda r: r["q_type"]),
    ("difficulty", 10, lambda r: r["difficulty"]),
    ("dimension", 18, lambda r: r["dimension"]),
    ("grounding", 10, lambda r: r["judge_scores"].get("grounding")),
    ("standalone", 11, lambda r: r["judge_scores"].get("standalone")),
    ("judge_rationale", 45, lambda r: " | ".join(x for x in (r["judge_scores"].get("grounding_rationale"),
                                                             r["judge_scores"].get("standalone_rationale")) if x)),
    ("paraphrase_1", 35, lambda r: (r["paraphrases"] + ["", ""])[0]),
    ("paraphrase_2", 35, lambda r: (r["paraphrases"] + ["", ""])[1]),
    ("original_question", 40, lambda r: r["repair"]["original_question"] if r["repair"] else ""),
    ("original_answer", 45, lambda r: r["repair"]["original_answer"] if r["repair"] else ""),
    ("repair_reason", 40, lambda r: r["repair"]["reason"] if r["repair"] else ""),
    ("origin", 22, lambda r: r.get("origin") or ""),
    ("qa_id", 40, lambda r: r["qa_id"]),
]


def citation_text(c):
    pages = c["pages"]
    if not pages:
        return c["title"]
    p = f"p. {pages[0]}" if len(pages) == 1 else f"pp. {pages[0]}-{pages[-1]}"
    return f"{c['title']}, {p}"


def col_letter(name):
    from openpyxl.utils import get_column_letter
    return get_column_letter([c[0] for c in COLUMNS].index(name) + 1)


def stratified_priority(rows, n_total, n_repaired, n_rejected, seed):
    """Pick priority rows: repaired-passed, rejected (repaired first), then stratified passed."""
    rng = random.Random(seed)

    def spread(pool, n, keys):
        """Round-robin over strata defined by keys, random within a stratum."""
        strata = defaultdict(list)
        for r in pool:
            strata[tuple(k(r) for k in keys)].append(r)
        for v in strata.values():
            rng.shuffle(v)
        order = sorted(strata, key=lambda s: (-len(strata[s]), s))
        out = []
        while len(out) < n and any(strata.values()):
            for s in order:
                if strata[s] and len(out) < n:
                    out.append(strata[s].pop())
        return out

    repaired_ok = [r for r in rows if r["repair"] and r["passed_filters"]]
    rejected = [r for r in rows if not r["passed_filters"]]
    pick = spread(repaired_ok, n_repaired, [lambda r: r["q_type"]])
    rej_rep = [r for r in rejected if r["repair"]]
    rej_plain = [r for r in rejected if not r["repair"]]
    half = (n_rejected + 1) // 2
    pick += rng.sample(rej_rep, min(half, len(rej_rep)))
    pick += rng.sample(rej_plain, min(n_rejected - min(half, len(rej_rep)), len(rej_plain)))
    chosen = {r["qa_id"] for r in pick}
    rest = [r for r in rows if r["passed_filters"] and not r["repair"] and r["qa_id"] not in chosen]
    pick += spread(rest, n_total - len(pick),
                   [lambda r: r["q_type"], lambda r: (r.get("origin") or "").split(":")[0],
                    lambda r: r["difficulty"] == "hard"])
    return pick


def build(rows, priority, out, run, counts):
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    ins.column_dimensions["A"].width = 110
    n_rep_ok = sum(bool(r["repair"]) and r["passed_filters"] for r in priority)
    n_rej = sum(not r["passed_filters"] for r in priority)
    n_rej_rep = sum(bool(r["repair"]) and not r["passed_filters"] for r in priority)
    n_other = len(priority) - n_rep_ok - n_rej
    lines = [
        (f"Pilot v2 QA review ({run})", True),
        ("", False),
        ("What to do", True),
        (f"1. Go to the Review sheet. Rows marked priority = TRUE (top {len(priority)}, highlighted) are a "
         f"stratified sample: {n_rep_ok} repaired pairs that passed, {n_rej} rejected pairs "
         f"({n_rej_rep} of them repaired) and {n_other} other passed pairs across question types, difficulty "
         f"and the new targeted chunks. Review those first; the other {len(rows) - len(priority)} rows are "
         "optional.", False),
        ("2. For each row, read the question, answer and evidence, then pick accept or reject in human_verdict "
         "(yellow column, dropdown).", False),
        ("3. If you reject, pick a reject_category (dropdown) and add a few words in human_notes.", False),
        ("4. The Agreement sheet updates automatically: your accept rate, how often you agree with the automatic "
         "filters, and your accept rate on repaired pairs.", False),
        ("", False),
        ("What is new in v2", True),
        ("repaired = TRUE: the first version failed grounding or standalone and was rewritten once by the generator, "
         "then re-judged. original_question / original_answer / repair_reason show the version that failed. "
         "The same judge passed the rewrite, so repaired pairs deserve a careful look.", False),
        ("answer no longer contains the citation; it is in the citation column (title and evidence pages). "
         "Judge the answer text only.", False),
        ("Paraphrases were checked and regenerated individually. If a paraphrase asks something different, note it "
         "in human_notes (the pair can still be accepted).", False),
        ("Answer length rule: 1-5 sentences for factual/definition, 2-5 for explanation/application. All rejected "
         "rows in this run failed only that rule.", False),
        ("", False),
        ("Accept a pair only if all three are true", True),
        ("Correct: the answer is fully supported by the evidence, with nothing added.", False),
        ("Standalone: the question makes sense to someone who has never seen the source.", False),
        ("Worth learning: it is about the domain (responsible AI, higher education), not trivia or study "
         "methodology.", False),
        ("", False),
        ("Reject categories", True),
        ("incorrect — answer contradicts or misreads the evidence", False),
        ("unsupported — answer adds claims the evidence does not make", False),
        ('not_standalone — question depends on unseen context ("the study", "this chapter")', False),
        ("low_value — trivial, methodology, bibliographic or appendix detail", False),
        ("unclear — question or answer is ambiguous or badly worded", False),
        ("other — explain in human_notes", False),
        ("", False),
        ("When done", True),
        ("Save this file (keep .xlsx), copy it into the project as data/qa_pairs/pilot_review_v2_reviewed.xlsx, "
         "and tell Claude Code to merge human_verdict, reject_category and human_notes by qa_id. Before saving, "
         "check that the Agreement sheet shows your rows as reviewed.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ins.cell(row=i, column=1, value=text or None)
        c.font = Font(bold=bold, size=13 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Review")
    for j, (name, width, _) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.fill = HEADER_FILL
        c.font = Font(bold=True, color="FFFFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 30
    prio_ids = {r["qa_id"] for r in priority}
    ordered = priority + [r for r in rows if r["qa_id"] not in prio_ids]
    for i, r in enumerate(ordered, 2):
        values = [i - 1, r["qa_id"] in prio_ids, None, None, None] + [g(r) for _, _, g in COLUMNS[5:]]
        for j, v in enumerate(values, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=3).fill = VERDICT_FILL
        ws.cell(row=i, column=4).fill = VERDICT_FILL
    last = len(ordered) + 1
    dv1 = DataValidation(type="list", formula1='"accept,reject"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1=f'"{",".join(CATEGORIES)}"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"C2:C{last}")
    dv2.add(f"D2:D{last}")
    ws.conditional_formatting.add(f"C2:C{last}", FormulaRule(formula=['$C2="accept"'],
                                  fill=PatternFill("solid", fgColor="FFC6EFCE")))
    ws.conditional_formatting.add(f"C2:C{last}", FormulaRule(formula=['$C2="reject"'],
                                  fill=PatternFill("solid", fgColor="FFFFC7CE")))
    ws.conditional_formatting.add(f"A2:B{last}", FormulaRule(formula=["$B2=TRUE"],
                                  fill=PatternFill("solid", fgColor="FFDDEBF7")))
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{col_letter('qa_id')}{last}"

    ag = wb.create_sheet("Agreement")
    ag.column_dimensions["A"].width = 48
    ag.column_dimensions["B"].width = 16
    ag.column_dimensions["C"].width = 16
    V, P, F, R, D = (f"Review!${col}$2:${col}${last}" for col in
                     ("C", "B", col_letter("passed_filters"), col_letter("repaired"), "D"))
    cells = [
        ("A1", "Review progress", True), ("A2", "Rows reviewed", False),
        ("B2", f'=COUNTIF({V},"accept")+COUNTIF({V},"reject")', False),
        ("A3", f"Priority rows reviewed (of {len(priority)})", False),
        ("B3", f'=COUNTIFS({P},TRUE,{V},"accept")+COUNTIFS({P},TRUE,{V},"reject")', False),
        ("A4", "Human accept rate", False), ("B4", f'=IFERROR(COUNTIF({V},"accept")/B2,"")', False),
        ("A5", "Human accept rate among filter-passed pairs", False),
        ("B5", f'=IFERROR(B9/(B9+C9),"")', False),
        ("A6", "Human accept rate among repaired pairs that passed", False),
        ("B6", f'=IFERROR(COUNTIFS({R},TRUE,{F},TRUE,{V},"accept")/(COUNTIFS({R},TRUE,{F},TRUE,{V},"accept")'
               f'+COUNTIFS({R},TRUE,{F},TRUE,{V},"reject")),"")', False),
        ("A8", "Confusion matrix (reviewed rows only)", True), ("B8", "Human accept", True),
        ("C8", "Human reject", True),
        ("A9", "Filters passed", False), ("B9", f'=COUNTIFS({F},TRUE,{V},"accept")', False),
        ("C9", f'=COUNTIFS({F},TRUE,{V},"reject")', False),
        ("A10", "Filters rejected", False), ("B10", f'=COUNTIFS({F},FALSE,{V},"accept")', False),
        ("C10", f'=COUNTIFS({F},FALSE,{V},"reject")', False),
        ("A12", "Observed agreement", False), ("B12", '=IFERROR((B9+C10)/(B9+C9+B10+C10),"")', False),
        ("A13", "Cohen's kappa", False),
        ("B13", '=IFERROR(((B9+C10)/(B9+C9+B10+C10)-((B9+C9)*(B9+B10)+(B10+C10)*(C9+C10))/(B9+C9+B10+C10)^2)'
                '/(1-((B9+C9)*(B9+B10)+(B10+C10)*(C9+C10))/(B9+C9+B10+C10)^2),"")', False),
        ("A15", "Human reject categories", True),
    ]
    for i, cat in enumerate(CATEGORIES):
        cells += [(f"A{16 + i}", cat, False), (f"B{16 + i}", f'=COUNTIF({D},"{cat}")', False)]
    cells.append((f"A{17 + len(CATEGORIES)}", "Rows 9-10 = automatic filter decision; columns B-C = your verdict. "
                  "Kappa above ~0.6 means the judge broadly matches your standard.", False))
    for coord, val, bold in cells:
        ag[coord] = val
        ag[coord].font = Font(bold=bold)
        ag[coord].alignment = Alignment(wrap_text=coord.startswith("A"), vertical="top")
    for coord in ("B4", "B5", "B6", "B12"):
        ag[coord].number_format = "0%"
    ag["B13"].number_format = "0.00"
    wb.save(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--run", default="pilot_v2")
    ap.add_argument("--out")
    ap.add_argument("--n-priority", type=int, default=40)
    ap.add_argument("--n-repaired", type=int, default=10)
    ap.add_argument("--n-rejected", type=int, default=4)
    ap.add_argument("--seed", type=int, default=7)
    args = ap.parse_args()
    cfg = load_config()
    rows = read_jsonl(run_paths(cfg, args.run)["final"])
    priority = stratified_priority(rows, args.n_priority, args.n_repaired, args.n_rejected, args.seed)
    out = repo_path(args.out) if args.out else repo_path(cfg["paths"]["qa_pairs"]) / \
        f"pilot_review_{args.run.split('_')[-1]}.xlsx"
    build(rows, priority, out, args.run, None)
    print(f"wrote {out}: {len(rows)} rows, {len(priority)} priority "
          f"({sum(bool(r['repair']) and r['passed_filters'] for r in priority)} repaired-passed, "
          f"{sum(not r['passed_filters'] for r in priority)} rejected)")


if __name__ == "__main__":
    main()
